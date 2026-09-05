from pathlib import Path
from openpyxl import load_workbook


BASE = Path(__file__).resolve().parent
BOOK = BASE / "modelo_paquetes_personal_2_etapas.xlsx"
REPORT = BASE / "RECOMENDACION_PAQUETES_PERSONAL.md"


def money(value):
    amount = float(value)
    return f"(${abs(amount):,.0f})" if amount < 0 else f"${amount:,.0f}"


def main():
    wb = load_workbook(BOOK, data_only=True)
    ws = wb["Personal"]
    secondments = wb["Secondments PU"]
    total = 20

    lines = [
        "# Modelo de personal en dos bloques",
        "",
        "**Fecha:** 5 de septiembre de 2026  ",
        "**Moneda:** dólares estadounidenses (USD)  ",
        "**Principio central:** el salario base no cambia y es idéntico a la tabla fuente.",
        "",
        "## Estructura correcta",
        "",
        "El modelo se divide en dos puntos de costo claramente separados. Durante **M1–M6**, las personas trabajan en rotación y reciben su salario base mensual sin cambios. El **único beneficio** en este período es el seguro médico. No se incluyen housing, home leave, transporte corporativo, impuestos pagados por la empresa, bonus, stock options ni vacaciones como costo adicional. JJI, por su condición Remote, no recibe seguro ni otros beneficios.",
        "",
        "Desde **M7**, se mantiene el mismo salario base y se incorporan los beneficios individuales definidos para expatriados, repatriados, personal local, Remote y On rotation.",
        "",
        "## Totales de los dos bloques",
        "",
        "| Punto de costo | Costo mensual | Costo de seis meses |",
        "|---|---:|---:|",
        f"| M1–M6: salario base + seguro médico únicamente | {money(ws['K20'].value)} | {money(ws['L20'].value)} |",
        f"| M7–M12: salario base + beneficios individuales | {money(ws['W20'].value)} | {money(ws['X20'].value)} |",
        "",
        f"El costo total del **Año 1 es {money(ws['Y20'].value)}**. Frente al paquete fuente anual de {money(ws['Z20'].value)}, el ahorro presupuestario es **{money(ws['AA20'].value)}**. El run-rate anual desde M13 es {money(ws['AB20'].value)}. Las stock options permanecen como TBD y no están incluidas en estos importes.",
        "",
        "## Secondments iniciales pagados por PU",
        "",
        "| Persona | Nombramiento inicial en PU | Modalidad | Pagador del salario | Salario base anual | Tratamiento en el total |",
        "|---|---|---|---|---:|---|",
    ]
    for row in range(11, 14):
        salary = secondments[f"G{row}"].value
        salary_text = "TBD" if salary == "TBD" else money(salary)
        lines.append(
            f"| {secondments[f'C{row}'].value} | {secondments[f'D{row}'].value} | {secondments[f'E{row}'].value} | "
            f"{secondments[f'F{row}'].value} | {salary_text} | {secondments[f'J{row}'].value} |"
        )
    lines.extend([
        "",
        "**PU paga el salario de los tres secondments.** Alexander Stulme y Alan McKeon ya están incluidos en el modelo principal con su salario base original, por lo que no se suman nuevamente. Se preserva la edición del usuario que establece para Martin Aguero un salario anual de **$180.000**; por ser una asignación adicional como Infrastructure Manager, se presenta separadamente y no se incluye en el total de las nueve posiciones. El costo salarial conocido de los tres secondments es $45.000 mensuales y $270.000 por seis meses.",
        "",
        "## Costo por persona",
        "",
        "| Persona | Salario base anual — sin cambios | Mensual M1–M6 | 6 meses M1–M6 | Condición M7+ | Mensual M7+ | 6 meses M7–M12 | Año 1 | Ahorro vs. fuente |",
        "|---|---:|---:|---:|---|---:|---:|---:|---:|",
    ])
    for row in range(11, 20):
        lines.append(
            f"| {ws[f'E{row}'].value} | {money(ws[f'F{row}'].value)} | {money(ws[f'K{row}'].value)} | "
            f"{money(ws[f'L{row}'].value)} | {ws[f'M{row}'].value} | {money(ws[f'W{row}'].value)} | "
            f"{money(ws[f'X{row}'].value)} | {money(ws[f'Y{row}'].value)} | {money(ws[f'AA{row}'].value)} |"
        )

    lines.extend([
        "",
        "## Beneficios desde M7",
        "",
        "| Persona | Housing | Seguro | Home leave | Transporte | Impuesto empresa | Bonus anual | Stock options | Vacaciones |",
        "|---|---:|---:|---:|---:|---:|---:|---|---:|",
    ])
    for row in range(11, 20):
        stock = ws[f"U{row}"].value
        stock_text = "TBD" if stock == "TBD" else money(stock or 0)
        lines.append(
            f"| {ws[f'E{row}'].value} | {money(ws[f'O{row}'].value)} | {money(ws[f'P{row}'].value)} | "
            f"{money(ws[f'Q{row}'].value)} | {money(ws[f'R{row}'].value)} | {money(ws[f'S{row}'].value)} | "
            f"{money(ws[f'T{row}'].value)} | {stock_text} | {int(ws[f'V{row}'].value or 0)} días |"
        )

    lines.extend([
        "",
        "## Confirmaciones individuales aplicadas",
        "",
        "**Juan Conde:** repatriado en Maracaibo; housing para allowance, staff house u hotel; seguro médico; transporte corporativo al trabajo; impuesto personal a su cargo; bonus de cuatro meses de salario; stock options TBD; 30 días de vacaciones.",
        "",
        "**Alexander Stulme:** repatriado en Maracaibo y vive en su propia casa; sin housing; seguro médico; transporte corporativo; impuesto personal a su cargo; bonus de cuatro meses; stock options TBD; 30 días de vacaciones.",
        "",
        "**Félix Valderrama:** mismo paquete de Juan Conde.",
        "",
        "**Alan McKeon:** expatriado con housing de $1.000 por mes desde M7.",
        "",
        "**JJI:** Remote con salario base únicamente y sin beneficios.",
        "",
        "**Jose Miguel y Operations Support:** se presupuestan como personal local en Maracaibo o alojado en staff house de PU. No reciben housing allowance en efectivo; cualquier costo operativo de la staff house deberá presupuestarse separadamente cuando esté disponible.",
        "",
        "## Base y advertencia",
        "",
        "**Base.** Los salarios anuales utilizados son exactamente $360.000, $216.000, $180.000, $216.000, $180.000, $120.000, $180.000, $48.000 y $48.000, iguales a la tabla fuente.",
        "",
        "> Este es un modelo presupuestario. La clasificación laboral, autorización de trabajo, impuestos, seguridad social, bonus, vacaciones y stock options requieren validación de asesores venezolanos y Recursos Humanos.",
        "",
    ])

    REPORT.write_text("\n".join(lines), encoding="utf-8")
    print(REPORT)


if __name__ == "__main__":
    main()
