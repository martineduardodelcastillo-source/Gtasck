from pathlib import Path
from math import isclose

from openpyxl import load_workbook


BASE = Path(__file__).resolve().parent
BOOK = BASE / "modelo_paquetes_personal_2_etapas.xlsx"
SOURCE = BASE / "source" / "package_source_2026-09-05.jpeg"

EXPECTED_BASE = {
    "Martin del Castillo": 360000,
    "Juan Conde": 216000,
    "Alexander Stulme": 180000,
    "Félix Valderrama": 216000,
    "Alan McKeon": 180000,
    "JJI": 120000,
    "Marcelo Dantas": 180000,
    "Jose Miguel": 48000,
    "TBD": 48000,
}

EXPECTED_MEDICAL_M1_M6 = {
    "Martin del Castillo": 12000,
    "Juan Conde": 12000,
    "Alexander Stulme": 12000,
    "Félix Valderrama": 12000,
    "Alan McKeon": 12000,
    "JJI": 0,
    "Marcelo Dantas": 12000,
    "Jose Miguel": 8000,
    "TBD": 8000,
}


def close(actual, expected, label, tolerance=0.05):
    assert actual is not None, f"{label}: valor ausente"
    assert isclose(float(actual), float(expected), rel_tol=0, abs_tol=tolerance), (
        f"{label}: esperado {expected}, obtenido {actual}"
    )


def main():
    assert BOOK.exists() and BOOK.stat().st_size > 0
    assert SOURCE.exists() and SOURCE.stat().st_size > 0

    formula_wb = load_workbook(BOOK, data_only=False)
    value_wb = load_workbook(BOOK, data_only=True)
    assert formula_wb.sheetnames == ["Resumen", "Personal", "Secondments PU", "Mensual", "Beneficios M7+", "Supuestos"]

    personal_f = formula_wb["Personal"]
    personal = value_wb["Personal"]
    rows = {}
    for row in range(1, personal.max_row + 1):
        value = personal[f"E{row}"].value
        if value in EXPECTED_BASE:
            rows[value] = row
    assert set(rows) == set(EXPECTED_BASE)

    # Salary is exactly the original salary for every person and never changes between blocks.
    for person, expected in EXPECTED_BASE.items():
        row = rows[person]
        close(personal[f"F{row}"].value, expected, f"{person} — salario base anual")
        close(personal[f"I{row}"].value, expected / 12, f"{person} — salario base mensual M1–M6")
        close(personal[f"G{row}"].value, EXPECTED_MEDICAL_M1_M6[person], f"{person} — seguro médico M1–M6")
        close(
            personal[f"K{row}"].value,
            expected / 12 + EXPECTED_MEDICAL_M1_M6[person] / 12,
            f"{person} — costo mensual M1–M6",
        )
        close(
            personal[f"L{row}"].value,
            (expected + EXPECTED_MEDICAL_M1_M6[person]) / 2,
            f"{person} — costo seis meses M1–M6",
        )
        assert personal[f"H{row}"].value == "Rotación"

    total_row = next(row for row in range(1, personal.max_row + 1) if personal[f"D{row}"].value == "TOTAL — 9 POSICIONES")
    close(personal[f"F{total_row}"].value, 1_548_000, "Total salario base anual")
    close(personal[f"G{total_row}"].value, 88_000, "Total seguro médico anual M1–M6")
    close(personal[f"K{total_row}"].value, 136_333.333333333, "Costo mensual M1–M6")
    close(personal[f"L{total_row}"].value, 818_000, "Bloque 1 — seis meses")
    close(personal[f"W{total_row}"].value, 189_700, "Costo mensual M7+")
    close(personal[f"X{total_row}"].value, 1_138_200, "Bloque 2 — seis meses")
    close(personal[f"Y{total_row}"].value, 1_956_200, "Costo total Año 1")
    close(personal[f"Z{total_row}"].value, 2_396_800, "Paquete fuente anual")
    close(personal[f"AA{total_row}"].value, 440_600, "Ahorro Año 1")
    close(personal[f"AB{total_row}"].value, 2_276_400, "Run-rate M13+")

    # Individual M7+ packages.
    checks = {
        "Juan Conde": ("Repatriado", "Maracaibo", 24000, 72000, "TBD", 30, 285000),
        "Alexander Stulme": ("Repatriado", "Maracaibo", 0, 60000, "TBD", 30, 231000),
        "Félix Valderrama": ("Repatriado", "Maracaibo", 24000, 72000, "TBD", 30, 285000),
        "Alan McKeon": ("Expatriado", "Maracaibo", 12000, 0, 0, 30, 231000),
        "JJI": ("Remote", "Remoto", 0, 0, 0, 30, 120000),
        "Jose Miguel": ("Local", "Maracaibo", 0, 0, 0, 30, 59600),
        "TBD": ("Local", "Maracaibo", 0, 0, 0, 30, 59600),
    }
    for person, (package, city, housing, bonus, stock, vacation, year1) in checks.items():
        row = rows[person]
        assert personal[f"M{row}"].value == package
        assert personal[f"N{row}"].value == city
        close(personal[f"O{row}"].value, housing, f"{person} — housing M7+")
        close(personal[f"T{row}"].value, bonus, f"{person} — bonus M7+")
        if stock == "TBD":
            assert personal[f"U{row}"].value == "TBD"
        else:
            close(personal[f"U{row}"].value, stock, f"{person} — stock options")
        close(personal[f"V{row}"].value, vacation, f"{person} — vacaciones")
        close(personal[f"Y{row}"].value, year1, f"{person} — Año 1")

    close(personal[f"V{total_row}"].value, 270, "Total días de vacaciones divulgados")

    # JJI has no benefits in either block.
    jji = rows["JJI"]
    for col in ("G", "O", "P", "Q", "R", "S", "T"):
        close(personal[f"{col}{jji}"].value, 0, f"JJI — {col}")

    # Monthly schedule must switch exactly at M7.
    monthly = value_wb["Mensual"]
    monthly_people = {}
    for row in range(1, monthly.max_row + 1):
        if monthly[f"C{row}"].value in EXPECTED_BASE:
            monthly_people[monthly[f"C{row}"].value] = row
    for person, source_row in rows.items():
        monthly_row = monthly_people[person]
        for col in range(5, 11):
            close(monthly.cell(monthly_row, col).value, personal[f"K{source_row}"].value, f"M1–M6 fila {monthly_row}")
        for col in range(11, 17):
            close(monthly.cell(monthly_row, col).value, personal[f"W{source_row}"].value, f"M7–M12 fila {monthly_row}")
        close(monthly[f"Q{monthly_row}"].value, personal[f"Y{source_row}"].value, f"Año 1 fila {monthly_row}")
    monthly_total = next(row for row in range(1, monthly.max_row + 1) if monthly[f"C{row}"].value == "TOTAL — 9 POSICIONES")
    close(monthly[f"E{monthly_total}"].value, 136_333.333333333, "Mensual total M1")
    close(monthly[f"K{monthly_total}"].value, 189_700, "Mensual total M7")
    close(monthly[f"Q{monthly_total}"].value, 1_956_200, "Mensual total Año 1")

    # Summary must reconcile to Personal.
    resumen = value_wb["Resumen"]
    summary_values = {resumen[f"C{row}"].value: resumen[f"D{row}"].value for row in range(1, resumen.max_row + 1)}
    close(summary_values["Costo mensual M1–M6"], personal[f"K{total_row}"].value, "Resumen mensual M1–M6")
    close(summary_values["Bloque 1 — seis meses"], personal[f"L{total_row}"].value, "Resumen bloque 1")
    close(summary_values["Costo mensual M7+"], personal[f"W{total_row}"].value, "Resumen mensual M7+")
    close(summary_values["Bloque 2 — seis meses"], personal[f"X{total_row}"].value, "Resumen bloque 2")
    close(summary_values["Costo total Año 1"], personal[f"Y{total_row}"].value, "Resumen Año 1")

    # PU pays the salaries of the three initial secondments; Martin Aguero is an additional cost.
    secondments = value_wb["Secondments PU"]
    secondment_rows = {secondments[f"C{row}"].value: row for row in range(1, secondments.max_row + 1) if secondments[f"C{row}"].value in {"Alexander Stulme", "Alan McKeon", "Martin Aguero"}}
    for person, appointment in (("Alexander Stulme", "General Manager — PU"), ("Alan McKeon", "Technical Manager — PU"), ("Martin Aguero", "Infrastructure Manager — PU")):
        row = secondment_rows[person]
        assert secondments[f"D{row}"].value == appointment
        assert secondments[f"E{row}"].value == "Secondment"
        assert secondments[f"F{row}"].value == "PU"
        close(secondments[f"G{row}"].value, 180_000, f"{person} — salario secondment")
        close(secondments[f"H{row}"].value, 15_000, f"{person} — mensual secondment")
        close(secondments[f"I{row}"].value, 90_000, f"{person} — seis meses secondment")
    second_total = next(row for row in range(1, secondments.max_row + 1) if secondments[f"C{row}"].value == "TOTAL CONOCIDO")
    close(secondments[f"G{second_total}"].value, 540_000, "Secondments — salario anual total")
    close(secondments[f"H{second_total}"].value, 45_000, "Secondments — salario mensual total")
    close(secondments[f"I{second_total}"].value, 270_000, "Secondments — seis meses total")

    formula_count = 0
    errors = []
    uncached = []
    blue_without_comments = []
    for sheet_name in formula_wb.sheetnames:
        fws = formula_wb[sheet_name]
        vws = value_wb[sheet_name]
        for row in fws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
                    cached = vws[cell.coordinate].value
                    if cached is None:
                        uncached.append(f"{sheet_name}!{cell.coordinate}")
                    elif isinstance(cached, str) and cached.startswith("#"):
                        errors.append(f"{sheet_name}!{cell.coordinate}={cached}")
                color = cell.font.color
                if color and color.type == "rgb" and color.rgb and color.rgb[-6:].upper() == "0000FF":
                    if cell.comment is None:
                        blue_without_comments.append(f"{sheet_name}!{cell.coordinate}")

    assert formula_count >= 430, f"Número inesperado de fórmulas: {formula_count}"
    assert not errors, errors
    assert not uncached, uncached[:20]
    assert not blue_without_comments, blue_without_comments[:20]

    print("VALIDATION PASSED")
    print("Base salary unchanged: $1,548,000 annual total")
    print("M1–M6 monthly: $136,333.33")
    print("Block 1 six months: $818,000")
    print("M7+ monthly: $189,700")
    print("Block 2 six months: $1,138,200")
    print("Year 1 total: $1,956,200")
    print("Year 1 savings vs. source: $440,600")
    print("Secondments paid by PU: Alexander, Alan and Martin Aguero")
    print("Martin Aguero salary preserved from user edit: $180,000")
    print(f"Formulas checked: {formula_count}")


if __name__ == "__main__":
    main()
