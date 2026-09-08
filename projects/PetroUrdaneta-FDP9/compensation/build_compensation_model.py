from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.datavalidation import DataValidation


BASE = Path(__file__).resolve().parent
OUT = BASE / "modelo_paquetes_personal_2_etapas.xlsx"
SOURCE_IMAGE = BASE / "source" / "package_source_2026-09-05.jpeg"

DARK_GREEN = "135B44"
LIGHT_GREEN = "CFE9E0"
WHITE = "FFFFFF"
BLACK = "000000"
BLUE = "0000FF"
GREEN = "008000"
RED = "C00000"
GRAY = "666666"
THIN = Side(style="thin", color="B7B7B7")
MEDIUM = Side(style="medium", color=DARK_GREEN)
DOUBLE = Side(style="double", color=BLACK)
CURR_FIRST = '$#,##0.0;($#,##0.0);-'
CURR = '#,##0.0;(#,##0.0);-'
INTEGER = '#,##0'

# Salary values are identical to the user's source table. No salary was changed.
PEOPLE = [
    # no, position, person, base, medical_b1, package_m7, city, housing, medical_m7,
    # home_leave, transport, employer_tax, bonus, stock, vacation, source_annual, notes
    (1, "COO / PU General Manager", "Martin del Castillo", 360000, 12000, "Expatriado", "Por definir", 36000, 12000, 20000, 24000, 72000, 0, 0, 30, 524000, "Expatriado; paquete internacional desde M7."),
    (2, "EPCM & Engineering Manager", "Juan Conde", 216000, 12000, "Repatriado", "Maracaibo", 24000, 12000, 0, 18000, 0, 72000, "TBD", 30, 325200, "Maracaibo: housing puede ser allowance, staff house u hotel. Impuesto personal a cargo del empleado."),
    (3, "Drilling & Well Services Manager", "Alexander Stulme", 180000, 12000, "Repatriado", "Maracaibo", 0, 12000, 0, 18000, 0, 60000, "TBD", 30, 282000, "Vive en su propia casa en Maracaibo; no recibe housing. Impuesto personal a su cargo."),
    (4, "Operations & Maintenance Manager", "Félix Valderrama", 216000, 12000, "Repatriado", "Maracaibo", 24000, 12000, 0, 18000, 0, 72000, "TBD", 30, 325200, "Mismo paquete de Juan Conde. Impuesto personal a cargo del empleado."),
    (5, "Technical Manager (Geosciences)", "Alan McKeon", 180000, 12000, "Expatriado", "Maracaibo", 12000, 12000, 12000, 18000, 36000, 0, 0, 30, 282000, "Housing expatriado de $1.000 por mes desde M7."),
    (6, "Reservoir Engineer", "JJI", 120000, 0, "Remote", "Remoto", 0, 0, 0, 0, 0, 0, 0, 30, 194000, "Remote: salario base únicamente; sin beneficios."),
    (7, "Rig Company Man", "Marcelo Dantas", 180000, 12000, "On rotation", "Por definir", 24000, 12000, 12000, 18000, 36000, 0, 0, 30, 282000, "On rotation: paquete provisional desde M7."),
    (8, "Planning & PMO", "Jose Miguel", 48000, 8000, "Local", "Maracaibo", 0, 8000, 0, 7200, 0, 0, 0, 30, 87200, "Local o staff house de PU; sin housing allowance en efectivo."),
    (9, "Operations Support", "TBD", 48000, 8000, "Local", "Maracaibo", 0, 8000, 0, 7200, 0, 0, 0, 30, 95200, "Local o staff house de PU; sin housing allowance en efectivo."),
]


def comment(text):
    return Comment(text, "Manus")


def source_note(row, field):
    return comment(
        f"Fuente: tabla de paquetes de personal suministrada por el usuario, 5 de septiembre de 2026. "
        f"Fila {row}, campo {field}. Archivo: {SOURCE_IMAGE.name}."
    )


def instruction_note(text):
    return comment(f"Fuente: instrucciones de la gerencia recibidas el 5 de septiembre de 2026. {text}")


def setup(ws, title, subtitle, units, end_col):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 20
    for col in range(3, end_col + 1):
        c = ws.cell(3, col)
        c.fill = PatternFill("solid", fgColor=DARK_GREEN)
        c.font = Font(name="Arial", size=16, bold=True, color=WHITE)
    ws.merge_cells(start_row=3, start_column=3, end_row=3, end_column=end_col)
    ws.cell(3, 3, title)
    ws.cell(5, 3, subtitle).font = Font(name="Arial", size=11, bold=True)
    ws.cell(6, 3, units).font = Font(name="Arial", size=10, italic=True, color=GRAY)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = f"{ws.title} | Página &P de &N"


def section(ws, row, title, start_col, end_col):
    for col in range(start_col, end_col + 1):
        c = ws.cell(row, col)
        c.fill = PatternFill("solid", fgColor=LIGHT_GREEN)
        c.font = Font(name="Arial", size=10, bold=True)
        c.border = Border(bottom=MEDIUM)
    ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    ws.cell(row, start_col, title)


def header(ws, row, start_col, headers):
    for offset, value in enumerate(headers):
        c = ws.cell(row, start_col + offset, value)
        c.fill = PatternFill("solid", fgColor=DARK_GREEN)
        c.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(top=MEDIUM, bottom=MEDIUM)
    ws.row_dimensions[row].height = 34


def input_cell(cell, value, note, fmt=None, align="right"):
    cell.value = value
    cell.font = Font(name="Arial", size=10, color=BLUE)
    cell.comment = note
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if fmt:
        cell.number_format = fmt


def formula_cell(cell, formula, fmt=None, cross=False, bold=False):
    cell.value = formula
    cell.font = Font(name="Arial", size=10, color=GREEN if cross else BLACK, bold=bold)
    cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    if fmt:
        cell.number_format = fmt


def borders(ws, start_row, end_row, start_col, end_col):
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            ws.cell(row, col).border = Border(
                top=THIN if row > start_row else Side(style=None),
                bottom=THIN if row < end_row else Side(style="thin", color=BLACK),
            )


def autofit(ws, start_col, end_col, start_row, end_row, maximum=28):
    for col in range(start_col, end_col + 1):
        max_len = 10
        for row in range(start_row, end_row + 1):
            value = ws.cell(row, col).value
            if value is not None and not (isinstance(value, str) and value.startswith("=")):
                max_len = max(max_len, max(len(x) for x in str(value).split("\n")))
        ws.column_dimensions[ws.cell(1, col).column_letter].width = min(max_len + 2, maximum)


def build_personal(wb):
    ws = wb.create_sheet("Personal")
    setup(
        ws,
        "Costo por persona — dos bloques",
        "Bloque 1: M1–M6 en rotación, salario base sin cambios + seguro médico únicamente | Bloque 2: M7–M12 con beneficios",
        "USD | Salario base anual idéntico a la tabla fuente | Azul = entrada; negro = fórmula",
        28,
    )
    section(ws, 8, "DOS PUNTOS DE COSTO POR PERSONA", 3, 28)
    headers = [
        "#", "Position", "Person", "Base Salary\nAnnual — unchanged", "Medical\nAnnual M1–M6", "Block 1\nStatus",
        "Base Salary\nMonthly M1–M6", "Medical\nMonthly M1–M6", "Total Monthly\nM1–M6", "Block 1\n6-Month Cost",
        "Package\nM7+", "City\nM7+", "Housing\nAnnual M7+", "Medical\nAnnual M7+", "Home Leave\nAnnual M7+",
        "Company Transport\nAnnual M7+", "Employer Tax\nAnnual M7+", "Annual Bonus\nM7+", "Stock Options\nAnnual Value", "Vacation\nDays",
        "Total Monthly\nM7+", "Block 2\n6-Month Cost", "Year 1\nTotal", "Source Annual\nPackage", "Year 1\nSavings", "Run-rate\nM13+",
    ]
    header(ws, 10, 3, headers)

    package_dv = DataValidation(type="list", formula1='"Expatriado,Repatriado,Remote,On rotation,Local,Por definir"')
    city_dv = DataValidation(type="list", formula1='"Caracas,Maracaibo,Otro,Remoto,Por definir"')
    ws.add_data_validation(package_dv)
    ws.add_data_validation(city_dv)

    for row, person_data in enumerate(PEOPLE, start=11):
        (number, position, person, base, medical_b1, package, city, housing, medical_m7, leave, transport, employer_tax, bonus, stock, vacation, source_annual, notes) = person_data
        input_cell(ws.cell(row, 3), number, source_note(number, "#"), INTEGER)
        input_cell(ws.cell(row, 4), position, source_note(number, "Position"), align="left")
        person_note = source_note(number, "Person") if person != "TBD" else instruction_note("Operations Support permanece como TBD.")
        input_cell(ws.cell(row, 5), person, person_note, align="left")
        input_cell(ws.cell(row, 6), base, source_note(number, "Base Salary Annual — sin cambios"), CURR_FIRST)
        med_note = source_note(number, "Medical Insurance") if medical_b1 else instruction_note("JJI es Remote y no recibe seguro médico.")
        input_cell(ws.cell(row, 7), medical_b1, med_note, CURR)
        input_cell(ws.cell(row, 8), "Rotación", instruction_note("M1–M6: sin beneficios, excepto seguro médico."), align="center")
        formula_cell(ws.cell(row, 9), f"=F{row}/12", CURR_FIRST)
        formula_cell(ws.cell(row, 10), f"=G{row}/12", CURR)
        formula_cell(ws.cell(row, 11), f"=SUM(I{row}:J{row})", CURR, bold=True)
        formula_cell(ws.cell(row, 12), f"=K{row}*6", CURR, bold=True)

        input_cell(ws.cell(row, 13), package, instruction_note(f"Condición desde M7: {package}."), align="center")
        input_cell(ws.cell(row, 14), city, instruction_note(f"Ciudad desde M7: {city}."), align="center")
        input_cell(ws.cell(row, 15), housing, instruction_note(f"Housing anual desde M7. {notes}"), CURR_FIRST)
        input_cell(ws.cell(row, 16), medical_m7, instruction_note("Seguro médico anual desde M7."), CURR)
        input_cell(ws.cell(row, 17), leave, instruction_note("Home leave anual desde M7; cero para repatriados y locales."), CURR)
        input_cell(ws.cell(row, 18), transport, instruction_note("Transporte provisto por la compañía para el trabajo / beneficio equivalente."), CURR)
        input_cell(ws.cell(row, 19), employer_tax, instruction_note("Costo fiscal asumido por la empresa; cero cuando el empleado asume su impuesto personal."), CURR)
        input_cell(ws.cell(row, 20), bonus, instruction_note("Bonus anual objetivo desde M7; cuatro meses de salario para Juan, Alexander y Félix."), CURR)
        stock_note = instruction_note("Stock options: valor pendiente de aprobación y valoración; no se incluye en el costo.")
        input_cell(ws.cell(row, 21), stock, stock_note, CURR if stock != "TBD" else None, align="right" if stock != "TBD" else "center")
        input_cell(ws.cell(row, 22), vacation, instruction_note("Se interpretó la instrucción como días de vacaciones pagadas; sin costo separado del salario."), INTEGER)
        formula_cell(ws.cell(row, 23), f"=SUM(F{row},O{row}:T{row})/12", CURR_FIRST, bold=True)
        formula_cell(ws.cell(row, 24), f"=W{row}*6", CURR, bold=True)
        formula_cell(ws.cell(row, 25), f"=L{row}+X{row}", CURR, bold=True)
        input_cell(ws.cell(row, 26), source_annual, source_note(number, "TOTAL Annual Package"), CURR)
        formula_cell(ws.cell(row, 27), f"=Z{row}-Y{row}", CURR)
        formula_cell(ws.cell(row, 28), f"=SUM(F{row},O{row}:T{row})", CURR, bold=True)
        package_dv.add(ws.cell(row, 13))
        city_dv.add(ws.cell(row, 14))
        ws.row_dimensions[row].height = 28

    total_row = 20
    ws.cell(total_row, 4, "TOTAL — 9 POSICIONES")
    ws.cell(total_row, 4).font = Font(name="Arial", size=10, bold=True)
    ws.merge_cells(start_row=total_row, start_column=4, end_row=total_row, end_column=5)
    for col in list(range(6, 13)) + list(range(15, 29)):
        letter = ws.cell(1, col).column_letter
        formula_cell(ws.cell(total_row, col), f"=SUM({letter}11:{letter}19)", CURR_FIRST if col in (6, 15, 23) else CURR, bold=True)
    for col in range(3, 29):
        ws.cell(total_row, col).border = Border(top=MEDIUM, bottom=DOUBLE)
    ws.row_dimensions[total_row].height = 24

    ws.auto_filter.ref = "C10:AB19"
    ws.freeze_panes = "I11"
    ws.print_title_rows = "10:10"
    ws.print_area = "B2:AB20"
    autofit(ws, 3, 28, 10, 20)
    ws.column_dimensions["D"].width = 34
    ws.column_dimensions["E"].width = 23
    ws.column_dimensions["U"].width = 18
    return ws


def build_resumen(wb):
    ws = wb.create_sheet("Resumen", 0)
    setup(
        ws,
        "Resumen ejecutivo — dos bloques de costo",
        "M1–M6: rotación con salario base + seguro médico únicamente | M7–M12: salario base + beneficios individuales",
        "USD | Stock options excluidas hasta valoración",
        12,
    )
    section(ws, 8, "INDICADORES PRINCIPALES", 3, 8)
    header(ws, 9, 3, ["Indicador", "Valor", "Definición", "Período", "Estado", "Fuente"])
    kpis = [
        ("Costo mensual M1–M6", "='Personal'!K20", "Salario base mensual + seguro médico únicamente", "M1–M6", "Calculado", "Personal"),
        ("Bloque 1 — seis meses", "='Personal'!L20", "Rotación; sin housing, home leave, transporte, impuesto empresarial ni bonus", "M1–M6", "Calculado", "Personal"),
        ("Costo mensual M7+", "='Personal'!W20", "Salario base + beneficios individuales", "M7+", "Calculado", "Personal"),
        ("Bloque 2 — seis meses", "='Personal'!X20", "Seis meses con paquetes individuales", "M7–M12", "Calculado", "Personal"),
        ("Costo total Año 1", "='Personal'!Y20", "Bloque 1 + Bloque 2", "Año 1", "Calculado", "Personal"),
        ("Paquete fuente anual", "='Personal'!Z20", "Referencia original de nueve posiciones", "12 meses", "Referencia", "Tabla fuente"),
        ("Ahorro Año 1 vs. fuente", "='Personal'!AA20", "Paquete fuente menos costo Año 1", "Año 1", "Calculado", "Personal"),
        ("Run-rate anual M13+", "='Personal'!AB20", "Costo recurrente anual con beneficios M7+", "M13+", "Calculado", "Personal"),
    ]
    for row, item in enumerate(kpis, start=10):
        ws.cell(row, 3, item[0]).font = Font(name="Arial", size=10, bold=row in (11, 14, 16, 17))
        formula_cell(ws.cell(row, 4), item[1], CURR_FIRST, cross=True, bold=row in (11, 14, 16, 17))
        for col, value in enumerate(item[2:], start=5):
            ws.cell(row, col, value)
            ws.cell(row, col).font = Font(name="Arial", size=9)
            ws.cell(row, col).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 28
    borders(ws, 10, 17, 3, 8)

    section(ws, 20, "DOS PUNTOS DE COSTO POR PERSONA", 3, 10)
    header(ws, 21, 3, ["Person", "Position", "Monthly M1–M6", "6M Cost M1–M6", "Package M7+", "Monthly M7+", "6M Cost M7–M12", "Year 1"])
    for target, source in enumerate(range(11, 20), start=22):
        refs = [f"='Personal'!E{source}", f"='Personal'!D{source}", f"='Personal'!K{source}", f"='Personal'!L{source}", f"='Personal'!M{source}", f"='Personal'!W{source}", f"='Personal'!X{source}", f"='Personal'!Y{source}"]
        for col, ref in enumerate(refs, start=3):
            formula_cell(ws.cell(target, col), ref, CURR_FIRST if col == 5 else CURR if col in (6, 8, 9, 10) else None, cross=True)
            if col in (3, 4, 7):
                ws.cell(target, col).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[target].height = 25
    borders(ws, 22, 30, 3, 10)

    section(ws, 33, "AHORRO POR PERSONA VS. PAQUETE FUENTE", 3, 9)
    header(ws, 34, 3, ["Person", "Source Annual Package", "Year 1 Cost", "Savings", "Primary driver", "Status", "Notes"])
    drivers = [
        "Seis meses sin beneficios salvo seguro médico.",
        "M1–M6 sin beneficios; M7 repatriado con bonus.",
        "M1–M6 sin beneficios; M7 sin housing.",
        "M1–M6 sin beneficios; M7 repatriado con bonus.",
        "M1–M6 sin beneficios; housing M7+ de $1.000/mes.",
        "Remote con salario base solamente.",
        "M1–M6 sin beneficios; paquete On rotation desde M7.",
        "M1–M6 salario + seguro; Local desde M7.",
        "M1–M6 salario + seguro; Local provisional desde M7.",
    ]
    for target, source in enumerate(range(11, 20), start=35):
        values = [f"='Personal'!E{source}", f"='Personal'!Z{source}", f"='Personal'!Y{source}", f"='Personal'!AA{source}"]
        for col, ref in enumerate(values, start=3):
            formula_cell(ws.cell(target, col), ref, CURR_FIRST if col == 4 else CURR if col in (5, 6) else None, cross=True)
            if col == 3:
                ws.cell(target, col).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(target, 7, drivers[target - 35])
        ws.cell(target, 8, "Confirmado" if source not in (11, 18, 19) else "Provisional")
        ws.cell(target, 9, "Stock options TBD" if source in (12, 13, 14) else "—")
        for col in range(7, 10):
            ws.cell(target, col).font = Font(name="Arial", size=9)
            ws.cell(target, col).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[target].height = 28
    borders(ws, 35, 43, 3, 9)

    chart = BarChart()
    chart.title = "Costo Año 1 por persona"
    chart.y_axis.title = "USD"
    chart.add_data(Reference(ws, min_col=5, min_row=34, max_row=43), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=3, min_row=35, max_row=43))
    chart.height = 8
    chart.width = 15
    chart.legend = None
    ws.add_chart(chart, "K20")

    section(ws, 46, "REGLA DEL MODELO", 3, 12)
    rules = [
        "1. El salario base anual no cambia entre bloques y coincide con la tabla fuente.",
        "2. M1–M6: el único beneficio es seguro médico; JJI no recibe beneficios por su condición Remote.",
        "3. M7–M12: se aplican los beneficios individuales documentados en Personal.",
        "4. Stock options permanecen en TBD y no se incluyen en costos; vacaciones no generan un costo separado del salario.",
    ]
    for row, text in enumerate(rules, start=47):
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=12)
        ws.cell(row, 3, text)
        ws.cell(row, 3).font = Font(name="Arial", size=10)
        ws.cell(row, 3).alignment = Alignment(wrap_text=True, vertical="center")

    section(ws, 53, "SECONDMENTS INICIALES — SALARIO PAGADO POR PU", 3, 10)
    header(ws, 54, 3, ["Person", "Initial PU Appointment", "Arrangement", "Salary Payer", "Annual Base Salary", "Monthly Salary", "6M Salary", "Treatment"])
    for target, source in enumerate(range(11, 14), start=55):
        source_cols = ("C", "D", "E", "F", "G", "H", "I", "J")
        for col, source_col in enumerate(source_cols, start=3):
            formula_cell(
                ws.cell(target, col),
                f"='Secondments PU'!{source_col}{source}",
                CURR_FIRST if col == 7 else CURR if col in (8, 9) else None,
                cross=True,
            )
            if col in (3, 4, 5, 6, 10):
                ws.cell(target, col).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[target].height = 28
    for col, source_col in enumerate(("C", "D", "E", "F", "G", "H", "I", "J"), start=3):
        formula_cell(
            ws.cell(58, col),
            f"='Secondments PU'!{source_col}15",
            CURR_FIRST if col == 7 else CURR if col in (8, 9) else None,
            cross=True,
            bold=True,
        )
    for col in range(3, 11):
        ws.cell(58, col).border = Border(top=MEDIUM, bottom=DOUBLE)

    ws.row_breaks.append(Break(id=52))
    ws.print_area = "B2:R58"
    ws.page_setup.fitToHeight = 0
    autofit(ws, 3, 10, 9, 43, maximum=30)
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["G"].width = 32
    return ws


def build_mensual(wb):
    ws = wb.create_sheet("Mensual")
    setup(
        ws,
        "Costo mensual por persona",
        "M1–M6: salario base + seguro médico | M7–M12: salario base + beneficios individuales",
        "USD | Costos recurrentes; stock options no valoradas",
        19,
    )
    section(ws, 8, "M1–M12 POR PERSONA", 3, 19)
    header(ws, 10, 3, ["Person", "Package M7+", "M1", "M2", "M3", "M4", "M5", "M6", "M7", "M8", "M9", "M10", "M11", "M12", "Year 1", "Monthly change M7 vs. M6", "Tax treatment M7+"])
    for target, source in enumerate(range(11, 20), start=11):
        formula_cell(ws.cell(target, 3), f"='Personal'!E{source}", cross=True)
        formula_cell(ws.cell(target, 4), f"='Personal'!M{source}", cross=True)
        ws.cell(target, 3).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(target, 4).alignment = Alignment(horizontal="left", vertical="center")
        for col in range(5, 11):
            formula_cell(ws.cell(target, col), f"='Personal'!K{source}", CURR_FIRST if col == 5 else CURR, cross=True)
        for col in range(11, 17):
            formula_cell(ws.cell(target, col), f"='Personal'!W{source}", CURR, cross=True)
        formula_cell(ws.cell(target, 17), f"=SUM(E{target}:P{target})", CURR, bold=True)
        formula_cell(ws.cell(target, 18), f"=K{target}-J{target}", CURR)
        formula_cell(ws.cell(target, 19), f'=IF(OR(D{target}="Expatriado",D{target}="On rotation"),"Empresa cubre impuesto",IF(OR(D{target}="Repatriado",D{target}="Local"),"Empleado asume impuesto",IF(D{target}="Remote","Sin beneficio fiscal","Definir")))')
        ws.cell(target, 19).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[target].height = 25
    total = 20
    ws.cell(total, 3, "TOTAL — 9 POSICIONES").font = Font(name="Arial", size=10, bold=True)
    ws.merge_cells(start_row=total, start_column=3, end_row=total, end_column=4)
    for col in range(5, 19):
        letter = ws.cell(1, col).column_letter
        formula_cell(ws.cell(total, col), f"=SUM({letter}11:{letter}19)", CURR_FIRST if col == 5 else CURR, bold=True)
    for col in range(3, 20):
        ws.cell(total, col).border = Border(top=MEDIUM, bottom=DOUBLE)
    ws.freeze_panes = "E11"
    ws.print_title_rows = "10:10"
    ws.print_area = "B2:S20"
    autofit(ws, 3, 19, 10, 20, maximum=24)
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["S"].width = 24
    return ws


def build_beneficios(wb):
    ws = wb.create_sheet("Beneficios M7+")
    setup(ws, "Beneficios individuales desde M7", "Detalle cualitativo y cuantitativo por persona", "USD anuales, salvo días", 13)
    section(ws, 8, "PAQUETE DE BENEFICIOS M7+", 3, 13)
    header(ws, 10, 3, ["Person", "Package", "City", "Housing", "Medical", "Home Leave", "Transport", "Employer Tax", "Bonus", "Stock Options", "Vacation Days"])
    for target, source in enumerate(range(11, 20), start=11):
        refs = ["E", "M", "N", "O", "P", "Q", "R", "S", "T", "U", "V"]
        for col, source_col in enumerate(refs, start=3):
            formula_cell(ws.cell(target, col), f"='Personal'!{source_col}{source}", CURR_FIRST if col == 6 else CURR if col in (7, 8, 9, 10, 11) else INTEGER if col == 13 else None, cross=True)
            if col in (3, 4, 5, 12):
                ws.cell(target, col).alignment = Alignment(horizontal="left" if col == 3 else "center", vertical="center", wrap_text=True)
        ws.row_dimensions[target].height = 24
    borders(ws, 11, 19, 3, 13)

    section(ws, 22, "DEFINICIONES", 3, 13)
    definitions = [
        "Repatriado: empleado asume su impuesto personal; housing y demás beneficios dependen del paquete individual.",
        "Expatriado: empresa cubre el impuesto presupuestado del país de trabajo y mantiene beneficios internacionales.",
        "Company Transport: transporte provisto por la compañía para el trabajo; no implica vehículo personal garantizado.",
        "Bonus: cuatro meses de salario anual para Juan, Alexander y Félix; se prorratea mensualmente desde M7.",
        "Stock options: beneficio reconocido, pero sin valor hasta aprobación y valoración del plan.",
        "Vacaciones: 30 días, según la edición directa del usuario; se consideran pagadas dentro del salario, sin costo incremental.",
    ]
    for row, text in enumerate(definitions, start=23):
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=13)
        ws.cell(row, 3, f"• {text}")
        ws.cell(row, 3).font = Font(name="Arial", size=10)
        ws.cell(row, 3).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 24
    ws.print_area = "B2:M28"
    ws.page_setup.fitToHeight = 1
    autofit(ws, 3, 13, 10, 19, maximum=24)
    ws.column_dimensions["C"].width = 24
    return ws


def build_secondments(wb):
    ws = wb.create_sheet("Secondments PU")
    setup(
        ws,
        "Secondments iniciales de PU",
        "Nombramientos iniciales cuyo salario es pagado por PU",
        "USD | Martin Aguero es un costo adicional separado del paquete de nueve posiciones",
        10,
    )
    section(ws, 8, "ASIGNACIONES INICIALES", 3, 10)
    header(ws, 10, 3, ["Person", "Initial PU Appointment", "Arrangement", "Salary Payer", "Annual Base Salary", "Monthly Salary", "6M Salary", "Treatment"])

    assignments = [
        ("Alexander Stulme", "General Manager — PU", "Secondment", "PU", "='Personal'!F13", "Incluido en el modelo de 9 posiciones"),
        ("Alan McKeon", "Technical Manager — PU", "Secondment", "PU", "='Personal'!F15", "Incluido en el modelo de 9 posiciones"),
        ("Martin Aguero", "Infrastructure Manager — PU", "Secondment", "PU", 180000, "Costo adicional de PU; separado del paquete de 9 posiciones"),
    ]
    for row, (person, appointment, arrangement, payer, salary, treatment) in enumerate(assignments, start=11):
        input_cell(ws.cell(row, 3), person, instruction_note(f"{person} fue indicado como secondment inicial de PU."), align="left")
        input_cell(ws.cell(row, 4), appointment, instruction_note(f"Nombramiento inicial: {appointment}."), align="left")
        input_cell(ws.cell(row, 5), arrangement, instruction_note("La asignación inicial es un secondment."), align="center")
        input_cell(ws.cell(row, 6), payer, instruction_note("PU paga el salario de la persona durante el secondment."), align="center")
        if salary == "TBD":
            input_cell(ws.cell(row, 7), salary, instruction_note("El salario base de Martin Aguero no fue suministrado; se excluye de los totales hasta recibirlo."), align="center")
            formula_cell(ws.cell(row, 8), f'=IF(ISNUMBER(G{row}),G{row}/12,"TBD")')
            formula_cell(ws.cell(row, 9), f'=IF(ISNUMBER(H{row}),H{row}*6,"TBD")')
        elif isinstance(salary, (int, float)):
            input_cell(ws.cell(row, 7), salary, instruction_note("Salario anual preservado de la edición directa del usuario del 5 de septiembre de 2026."), CURR_FIRST)
            formula_cell(ws.cell(row, 8), f"=G{row}/12", CURR)
            formula_cell(ws.cell(row, 9), f"=H{row}*6", CURR)
        else:
            formula_cell(ws.cell(row, 7), salary, CURR_FIRST, cross=True)
            formula_cell(ws.cell(row, 8), f"=G{row}/12", CURR)
            formula_cell(ws.cell(row, 9), f"=H{row}*6", CURR)
        ws.cell(row, 10, treatment)
        ws.cell(row, 10).font = Font(name="Arial", size=9)
        ws.cell(row, 10).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 28
    borders(ws, 11, 13, 3, 10)

    ws.cell(15, 3, "TOTAL CONOCIDO")
    ws.cell(15, 3).font = Font(name="Arial", size=10, bold=True)
    ws.merge_cells("C15:F15")
    formula_cell(ws.cell(15, 7), "=SUM(G11:G13)", CURR_FIRST, bold=True)
    formula_cell(ws.cell(15, 8), "=SUM(H11:H13)", CURR, bold=True)
    formula_cell(ws.cell(15, 9), "=SUM(I11:I13)", CURR, bold=True)
    ws.cell(15, 10, "Incluye Martin Aguero; no duplicar Alexander y Alan")
    ws.cell(15, 10).font = Font(name="Arial", size=10, bold=True, color=RED)
    for col in range(3, 11):
        ws.cell(15, col).border = Border(top=MEDIUM, bottom=DOUBLE)

    section(ws, 18, "TRATAMIENTO EN EL MODELO", 3, 10)
    notes = [
        "Alexander Stulme y Alan McKeon ya están incluidos en el modelo principal con sus salarios base originales sin cambios.",
        "Martin Aguero es una asignación adicional con salario anual de $180.000; se presenta separado del costo de nueve posiciones.",
        "PU se identifica como pagador del salario para los tres secondments.",
    ]
    for row, text in enumerate(notes, start=19):
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=10)
        ws.cell(row, 3, f"• {text}")
        ws.cell(row, 3).font = Font(name="Arial", size=10)
        ws.cell(row, 3).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 25
    ws.print_area = "B2:J21"
    ws.page_setup.fitToHeight = 1
    autofit(ws, 3, 10, 10, 15, maximum=32)
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["J"].width = 35
    return ws


def build_supuestos(wb):
    ws = wb.create_sheet("Supuestos")
    setup(ws, "Supuestos del modelo reconstruido", "Reglas acordadas para los dos bloques", "Celdas azules = supuestos editables", 10)
    section(ws, 8, "REGLAS PRINCIPALES", 3, 10)
    rules = [
        ("Meses Bloque 1", 6, "M1–M6"),
        ("Meses Bloque 2", 6, "M7–M12"),
        ("Beneficio permitido en Bloque 1", "Seguro médico únicamente", "JJI no recibe beneficios por ser Remote"),
        ("Salario base", "Sin cambios", "Misma base anual de la tabla fuente"),
        ("Stock options", "TBD", "Excluidas de costo"),
        ("Vacaciones", "30 días", "Edición del usuario preservada; sin costo incremental"),
    ]
    header(ws, 10, 3, ["Supuesto", "Valor", "Aplicación"])
    for row, (label, value, application) in enumerate(rules, start=11):
        ws.cell(row, 3, label).font = Font(name="Arial", size=10)
        input_cell(ws.cell(row, 4), value, instruction_note(application), INTEGER if isinstance(value, int) else None, align="right" if isinstance(value, int) else "left")
        ws.cell(row, 5, application).font = Font(name="Arial", size=10)
        ws.cell(row, 5).alignment = Alignment(wrap_text=True, vertical="center")
    borders(ws, 11, 16, 3, 5)

    section(ws, 19, "ADVERTENCIA", 3, 10)
    warning = (
        "El modelo es presupuestario. La autorización de trabajo, clasificación laboral, residencia fiscal, impuesto personal, "
        "cargas patronales, bonus, vacaciones y stock options deben ser validados por asesores venezolanos y Recursos Humanos."
    )
    ws.merge_cells("C20:J22")
    ws["C20"] = warning
    ws["C20"].font = Font(name="Arial", size=10, bold=True, color=RED)
    ws["C20"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.print_area = "B2:J22"
    ws.page_setup.fitToHeight = 1
    autofit(ws, 3, 10, 10, 16, maximum=35)
    return ws


def main():
    assert SOURCE_IMAGE.exists()
    assert sum(p[3] for p in PEOPLE) == 1_548_000
    assert sum(p[15] for p in PEOPLE) == 2_396_800
    wb = Workbook()
    wb.remove(wb.active)
    build_personal(wb)
    build_resumen(wb)
    build_secondments(wb)
    build_mensual(wb)
    build_beneficios(wb)
    build_supuestos(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
