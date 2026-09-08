from pathlib import Path
from openpyxl import load_workbook

book = Path(__file__).resolve().parent / "personnel_base_salary_contract_types.xlsx"
wb = load_workbook(book, data_only=True)

for sheet_name, cells in {
    "Monthly Costs": ["Q17", "R17", "T17", "W17", "X17", "Y12", "Z12"],
    "PU Secondments": ["F15", "L15", "N15", "Q15", "R15", "D20", "D21", "D22", "J20", "J21", "J22", "P20", "P21", "P22"],
    "Alan Expat Package": ["D9", "D10", "D11", "D12", "D13", "D14", "D19", "D20", "D21"],
}.items():
    ws = wb[sheet_name]
    print(f"{sheet_name}:")
    for cell in cells:
        print(f"  {cell}={ws[cell].value}")

monthly = wb["Monthly Costs"]
secondments = wb["PU Secondments"]
assert monthly["Q17"].value == 108333.333333333
assert monthly["R17"].value == 184833.333333333
assert monthly["T17"].value == 114333.333333333
assert monthly["W17"].value == 114333.333333333
assert monthly["X17"].value == 1376500
assert secondments["D22"].value == 59333.3333333333
assert secondments["J22"].value == 113333.333333333
assert secondments["P22"].value == 766000
assert wb["Alan Expat Package"]["D21"].value == 24000
print("CALCULATED VALUES VALIDATED")
