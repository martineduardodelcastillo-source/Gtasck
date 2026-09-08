from pathlib import Path
from openpyxl import load_workbook

book = Path(__file__).resolve().parent / "manpower_cost_summary_single_sheet.xlsx"
wb = load_workbook(book, data_only=True)
ws = wb["Manpower Cost Summary"]

checks = {
    "gross_m1": ws["D13"].value,
    "pu_m1_deduction": ws["D14"].value,
    "net_m1": ws["D15"].value,
    "gross_m7": ws["J13"].value,
    "pu_m7_deduction": ws["J14"].value,
    "net_m7": ws["J15"].value,
    "gross_m9": ws["L13"].value,
    "pu_m9_deduction": ws["L14"].value,
    "net_m9": ws["L15"].value,
    "gross_year1": ws["P13"].value,
    "pu_year1_deduction": ws["P14"].value,
    "net_year1": ws["P15"].value,
    "alan_housing": ws["N32"].value,
    "alan_home_leave": ws["O32"].value,
    "martin_health_care": ws["L37"].value,
    "martin_local_vacation": ws["P37"].value,
    "martin_year1": ws["Q51"].value,
}
for label, value in checks.items():
    print(f"{label}={value}")

assert checks["gross_m1"] == 122333.333333333
assert checks["pu_m1_deduction"] == -48000
assert checks["net_m1"] == 74333.3333333333
assert checks["gross_m7"] == 200833.333333333
assert checks["pu_m7_deduction"] == -72500
assert checks["net_m7"] == 128333.333333333
assert checks["gross_m9"] == 124333.333333333
assert checks["pu_m9_deduction"] == -50000
assert checks["net_m9"] == 74333.3333333333
assert checks["gross_year1"] == 1556500
assert checks["pu_year1_deduction"] == -610500
assert checks["net_year1"] == 946000
assert checks["alan_housing"] == 2000
assert checks["alan_home_leave"] == "Home leave"
assert checks["martin_health_care"] == 12000
assert checks["martin_local_vacation"] == 30
assert checks["martin_year1"] == 192000
print("CALCULATED VALUES VALIDATED")
