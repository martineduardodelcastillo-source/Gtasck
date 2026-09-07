from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


BASE = Path(__file__).resolve().parent
OUT = BASE / "personnel_base_salary_contract_types.xlsx"
SOURCE_IMAGE = Path("/home/ubuntu/upload/pasted_file_yfFQ2v_image.png")
SOURCE_DATE = "7 September 2026"

DARK_GREEN = "135B44"
LIGHT_GREEN = "CFE9E0"
LIGHT_GRAY = "E7E5E4"
WHITE = "FFFFFF"
BLUE = "0000FF"
BLACK = "000000"
GREEN = "008000"
GRID = "B7B7B7"
WARNING = "FFF2CC"
CURRENCY_FIRST = '$#,##0.0;[Red]($#,##0.0);-'
CURRENCY = '#,##0.0;[Red](#,##0.0);-'

PEOPLE = [
    # number, position, person, annual base salary, contract type, work arrangement, status, notes
    (1, "COO / PU General Manager", "Martin del Castillo", 360000, "Expat", "Fixed", "Confirmed", "Expat classification carried forward from prior instructions."),
    (2, "EPCM & Engineering Manager", "Juan Conde", 216000, "Repatriate", "Fixed", "Confirmed", "Repatriate; employee assumes personal Venezuelan income tax from M7."),
    (3, "Drilling & Well Services Manager", "Alexander Stulme", 180000, "Repatriate", "Fixed", "Confirmed", "Repatriate; initially appointed General Manager of PU on secondment."),
    (4, "Operations & Maintenance Manager", "Félix Valderrama", 216000, "Repatriate", "Fixed", "Confirmed", "Same contract family as Juan Conde."),
    (5, "Technical Manager (Geosciences)", "Alan McKeon", 180000, "Expat", "Fixed", "Confirmed", "Expat; initially appointed Technical Manager of PU on secondment."),
    (6, "Reservoir Engineer", "JJI", 120000, "Local", "Fixed", "Assumption", "Remote role. Local contract is an editable assumption pending confirmation."),
    (7, "Rig Company Man", "Marcelo Dantas", 180000, "Expat", "Rotation", "Assumption", "Rotation confirmed from prior instructions; Expat contract is an editable assumption."),
    (8, "Planning & PMO", "Jose Miguel", 48000, "Local", "Fixed", "Confirmed", "Local in Maracaibo or staff house; no cash housing allowance by default."),
    (9, "Operations Support", "Leticia Almeida", 48000, "Local", "Fixed", "Confirmed", "Local in Maracaibo or staff house; no cash housing allowance by default."),
]


def source_comment(field, extra=""):
    base = (
        f"Source: compensation table image {SOURCE_IMAGE.name}, provided {SOURCE_DATE}; "
        f"field: {field}."
    )
    return Comment(f"{base} {extra}".strip(), "Manus")


def instruction_comment(text):
    return Comment(f"Source: management instructions in this task, through {SOURCE_DATE}. {text}", "Manus")


def set_input(cell, value, comment):
    cell.value = value
    cell.font = Font(name="Arial", size=10, color=BLUE)
    cell.comment = comment


def set_formula(cell, formula, number_format=None, cross_sheet=False):
    cell.value = formula
    cell.font = Font(name="Arial", size=10, color=GREEN if cross_sheet else BLACK)
    if number_format:
        cell.number_format = number_format


def title_block(ws, title, subtitle, units, end_col):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 20
    ws.merge_cells(start_row=3, start_column=3, end_row=3, end_column=end_col)
    ws.cell(3, 3, title)
    ws.cell(3, 3).fill = PatternFill("solid", fgColor=DARK_GREEN)
    ws.cell(3, 3).font = Font(name="Arial", size=16, bold=True, color=WHITE)
    ws.cell(3, 3).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 25
    ws.merge_cells(start_row=5, start_column=3, end_row=5, end_column=end_col)
    ws.cell(5, 3, subtitle)
    ws.cell(5, 3).font = Font(name="Arial", size=11, bold=True, color=BLACK)
    ws.merge_cells(start_row=6, start_column=3, end_row=6, end_column=end_col)
    ws.cell(6, 3, units)
    ws.cell(6, 3).font = Font(name="Arial", size=9, italic=True, color="666666")


def header_row(ws, row, start_col, headers):
    for col, header in enumerate(headers, start=start_col):
        cell = ws.cell(row, col, header)
        cell.fill = PatternFill("solid", fgColor=LIGHT_GREEN)
        cell.font = Font(name="Arial", size=10, bold=True, color=BLACK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=DARK_GREEN))
    ws.row_dimensions[row].height = 32


def apply_borders(ws, start_row, end_row, start_col, end_col):
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row, col)
            cell.border = Border(
                left=Side(style="thin", color=GRID) if col == start_col else Side(style=None),
                right=Side(style="thin", color=GRID) if col == end_col else Side(style=None),
                top=Side(style="thin", color=GRID),
                bottom=Side(style="thin", color=GRID),
            )


def autofit(ws, start_col, end_col, start_row, end_row):
    for col in range(start_col, end_col + 1):
        max_len = 0
        for row in range(start_row, end_row + 1):
            value = ws.cell(row, col).value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 3, 12), 38)


def build_personnel(wb):
    ws = wb.active
    ws.title = "Personnel"
    title_block(
        ws,
        "Personnel — Base Salary and Contract Classification",
        "Three contract types: Expat, Repatriate and Local | Work arrangement shown separately as Fixed or Rotation",
        "USD | Annual base salaries transcribed from the source image | Blue = editable input; black = formula",
        12,
    )

    headers = [
        "#",
        "Position",
        "Person",
        "Base Salary Annual",
        "Base Salary Monthly",
        "Contract Type",
        "Expatriate?",
        "Work Arrangement",
        "Classification Status",
        "Notes",
    ]
    header_row(ws, 8, 3, headers)

    for row, record in enumerate(PEOPLE, start=9):
        number, position, person, annual_salary, contract_type, arrangement, status, notes = record
        set_input(ws.cell(row, 3), number, source_comment("row number"))
        set_input(ws.cell(row, 4), position, source_comment("Position"))
        set_input(ws.cell(row, 5), person, source_comment("Person"))
        set_input(ws.cell(row, 6), annual_salary, source_comment("Base Salary (Annual)"))
        ws.cell(row, 6).number_format = CURRENCY_FIRST if row == 9 else CURRENCY
        set_formula(ws.cell(row, 7), f"=F{row}/12", CURRENCY_FIRST if row == 9 else CURRENCY)
        set_input(
            ws.cell(row, 8),
            contract_type,
            instruction_comment("Contract type must be Expat, Repatriate or Local. Edit if management confirms a different classification."),
        )
        set_formula(ws.cell(row, 9), f'=IF(H{row}="Expat","Yes","No")')
        set_input(
            ws.cell(row, 10),
            arrangement,
            instruction_comment("Work arrangement must be Fixed or Rotation. Remote is documented in Notes, not treated as a fourth contract type."),
        )
        set_input(
            ws.cell(row, 11),
            status,
            instruction_comment("Confirmed means directly established in prior instructions; Assumption remains editable and requires confirmation."),
        )
        set_input(ws.cell(row, 12), notes, instruction_comment(notes))
        for col in range(3, 13):
            ws.cell(row, col).alignment = Alignment(
                horizontal="right" if col in (3, 6, 7) else "center" if col in (8, 9, 10, 11) else "left",
                vertical="center",
                wrap_text=col in (4, 12),
            )
        ws.row_dimensions[row].height = 30

    total_row = 18
    ws.merge_cells(start_row=total_row, start_column=3, end_row=total_row, end_column=5)
    ws.cell(total_row, 3, "TOTAL — 9 PEOPLE")
    ws.cell(total_row, 3).font = Font(name="Arial", size=10, bold=True, color=BLACK)
    set_formula(ws.cell(total_row, 6), "=SUM(F9:F17)", CURRENCY_FIRST)
    set_formula(ws.cell(total_row, 7), "=SUM(G9:G17)", CURRENCY_FIRST)
    for col in range(3, 13):
        ws.cell(total_row, col).border = Border(top=Side(style="medium", color=BLACK), bottom=Side(style="double", color=BLACK))
        if col >= 6:
            ws.cell(total_row, col).alignment = Alignment(horizontal="right", vertical="center")

    # Summary counts are formulas so they respond to user edits.
    ws.merge_cells("C21:L21")
    ws["C21"] = "CLASSIFICATION SUMMARY"
    ws["C21"].fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    ws["C21"].font = Font(name="Arial", size=10, bold=True, color=BLACK)
    header_row(ws, 22, 3, ["Metric", "Count", "Meaning"])
    summary = [
        ("Expat", '=COUNTIF(H9:H17,"Expat")', "Contract type Expat"),
        ("Repatriate", '=COUNTIF(H9:H17,"Repatriate")', "Contract type Repatriate"),
        ("Local", '=COUNTIF(H9:H17,"Local")', "Contract type Local"),
        ("Fixed", '=COUNTIF(J9:J17,"Fixed")', "Fixed work arrangement"),
        ("Rotation", '=COUNTIF(J9:J17,"Rotation")', "Rotation work arrangement"),
        ("Assumptions to confirm", '=COUNTIF(K9:K17,"Assumption")', "Editable classifications requiring confirmation"),
    ]
    for row, (metric, formula, meaning) in enumerate(summary, start=23):
        ws.cell(row, 3, metric).font = Font(name="Arial", size=10, color=BLACK)
        set_formula(ws.cell(row, 4), formula, '#,##0')
        ws.cell(row, 5, meaning).font = Font(name="Arial", size=10, color=BLACK)
    apply_borders(ws, 22, 28, 3, 5)

    contract_validation = DataValidation(type="list", formula1='"Expat,Repatriate,Local"', allow_blank=False)
    arrangement_validation = DataValidation(type="list", formula1='"Fixed,Rotation"', allow_blank=False)
    status_validation = DataValidation(type="list", formula1='"Confirmed,Assumption"', allow_blank=False)
    ws.add_data_validation(contract_validation)
    ws.add_data_validation(arrangement_validation)
    ws.add_data_validation(status_validation)
    contract_validation.add("H9:H17")
    arrangement_validation.add("J9:J17")
    status_validation.add("K9:K17")

    ws.conditional_formatting.add(
        "K9:K17",
        FormulaRule(formula=['K9="Assumption"'], fill=PatternFill("solid", fgColor=WARNING)),
    )
    apply_borders(ws, 8, 17, 3, 12)
    ws.freeze_panes = "F9"
    autofit(ws, 3, 12, 8, 28)
    ws.column_dimensions["D"].width = max(ws.column_dimensions["D"].width, 38)
    ws.column_dimensions["E"].width = max(ws.column_dimensions["E"].width, 24)
    ws.column_dimensions["L"].width = 38
    ws.print_area = "B2:L28"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Personnel | Page &P of &N"
    ws.oddFooter.center.size = 8
    return ws


def build_rules(wb):
    ws = wb.create_sheet("Contract Rules")
    title_block(
        ws,
        "Contract Benefit Rules by Stage",
        "Default policy view — individual packages may override these rules",
        "M1–M6 and M7+ | Text rules are editable assumptions, not legal or tax advice",
        11,
    )

    headers = [
        "Contract Type",
        "M1–M6 Benefits",
        "M7+ Medical",
        "M7+ Housing",
        "M7+ Tax Treatment",
        "M7+ Other Benefits",
        "Default Status",
        "Notes",
        "People Count",
    ]
    header_row(ws, 8, 3, headers)
    rules = [
        (
            "Expat",
            "Medical insurance only",
            "Company-provided medical insurance",
            "Housing / allowance according to approved package",
            "Company covers or equalizes work-country tax",
            "Home leave and company transport; other benefits by approved package",
            "Provisional",
            "Fixed and Rotation arrangements may have different housing and travel mechanics.",
        ),
        (
            "Repatriate",
            "Medical insurance only",
            "Company-provided medical insurance",
            "Depends on city: allowance, staff house, hotel or no housing",
            "Employee assumes personal Venezuelan income tax",
            "Company transport; bonus, stock options and vacation only when individually approved",
            "Provisional",
            "Juan, Alexander and Félix have individual terms that should be documented separately.",
        ),
        (
            "Local",
            "Medical insurance only",
            "Medical insurance according to local package",
            "Local residence or staff house; no cash allowance by default",
            "Employee assumes personal Venezuelan income tax",
            "Company transport according to role; no international home leave by default",
            "Provisional",
            "Jose Miguel and Operations Support are assumed Local / staff house in Maracaibo.",
        ),
    ]
    for row, record in enumerate(rules, start=9):
        for col, value in enumerate(record, start=3):
            set_input(ws.cell(row, col), value, instruction_comment(f"Default contract rule: {value}"))
            ws.cell(row, col).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        set_formula(ws.cell(row, 11), f'=COUNTIF(Personnel!H9:H17,C{row})', '#,##0', cross_sheet=True)
        ws.cell(row, 11).alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[row].height = 60
    apply_borders(ws, 8, 11, 3, 11)

    ws.merge_cells("C14:K14")
    ws["C14"] = "WORK ARRANGEMENT RULES"
    ws["C14"].fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    ws["C14"].font = Font(name="Arial", size=10, bold=True, color=BLACK)
    header_row(ws, 15, 3, ["Work Arrangement", "Application", "Housing / Lodging", "Travel / Home Leave", "Notes", "People Count"])
    arrangements = [
        ("Fixed", "Continuous assignment to the role", "Based on contract type, city and approved package", "According to contract type", "JJI is Fixed but Remote; Remote is noted separately.", '=COUNTIF(Personnel!J9:J17,C16)'),
        ("Rotation", "Roster-based presence in Venezuela", "Hotel, staff house or rotational lodging; avoid duplicate permanent housing", "Rotation travel schedule replaces or supplements standard home leave", "Marcelo Dantas is currently shown as Rotation.", '=COUNTIF(Personnel!J9:J17,C17)'),
    ]
    for row, record in enumerate(arrangements, start=16):
        for col, value in enumerate(record[:5], start=3):
            set_input(ws.cell(row, col), value, instruction_comment(f"Default work-arrangement rule: {value}"))
            ws.cell(row, col).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        set_formula(ws.cell(row, 8), record[5], '#,##0', cross_sheet=True)
        ws.cell(row, 8).alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[row].height = 55
    apply_borders(ws, 15, 17, 3, 8)

    ws.merge_cells("C20:K22")
    ws["C20"] = (
        "Important: M1–M6 is modeled as the initial stage with base salary plus medical insurance only. "
        "From M7, benefits change according to contract type and work arrangement. Tax, immigration, employment status, "
        "housing and benefits must be confirmed by Venezuelan legal, tax and HR advisers before implementation."
    )
    ws["C20"].font = Font(name="Arial", size=10, bold=True, color="C00000")
    ws["C20"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    autofit(ws, 3, 11, 8, 22)
    for col in range(3, 11):
        ws.column_dimensions[get_column_letter(col)].width = min(max(ws.column_dimensions[get_column_letter(col)].width, 20), 32)
    ws.column_dimensions["J"].width = 36
    ws.column_dimensions["K"].width = 14
    ws.print_area = "B2:K22"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Contract Rules | Page &P of &N"
    ws.oddFooter.center.size = 8
    return ws


def build():
    if not SOURCE_IMAGE.exists():
        raise FileNotFoundError(SOURCE_IMAGE)
    wb = Workbook()
    build_personnel(wb)
    build_rules(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUT)
    print(OUT)


def validate():
    wb = load_workbook(OUT, data_only=False)
    assert wb.sheetnames == ["Personnel", "Contract Rules"]
    ws = wb["Personnel"]
    expected_names = [record[2] for record in PEOPLE]
    expected_salaries = [record[3] for record in PEOPLE]
    assert [ws[f"E{row}"].value for row in range(9, 18)] == expected_names
    assert [ws[f"F{row}"].value for row in range(9, 18)] == expected_salaries
    assert sum(expected_salaries) == 1_548_000
    assert ws["G9"].value == "=F9/12"
    assert ws["I9"].value == '=IF(H9="Expat","Yes","No")'
    assert [ws[f"H{row}"].value for row in range(9, 18)].count("Expat") == 3
    assert [ws[f"H{row}"].value for row in range(9, 18)].count("Repatriate") == 3
    assert [ws[f"H{row}"].value for row in range(9, 18)].count("Local") == 3
    assert [ws[f"J{row}"].value for row in range(9, 18)].count("Rotation") == 1
    assert [ws[f"J{row}"].value for row in range(9, 18)].count("Fixed") == 8
    rules = wb["Contract Rules"]
    assert [rules[f"C{row}"].value for row in range(9, 12)] == ["Expat", "Repatriate", "Local"]
    print("VALIDATED: 9 people | $1,548,000 annual base salary | 3 contract types | 2 work arrangements")


if __name__ == "__main__":
    build()
    validate()
