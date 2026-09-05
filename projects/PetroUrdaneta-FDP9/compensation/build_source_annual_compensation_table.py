from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "source_annual_compensation_table.xlsx"
SOURCE = Path("/home/ubuntu/upload/pasted_file_aWe6xG_image.png")

HEADERS = [
    "#",
    "Position",
    "Person",
    "Base Salary (Annual)",
    "Housing (Annual)",
    "Medical Insurance",
    "Home Leave (2 trips/yr)",
    "Company Vehicle",
    "Tax & Social",
    "TOTAL Annual Package",
    "Secondment Status",
]

ROWS = [
    (1, "COO / PU General Manager", "Martin del Castillo", 360000, 36000, 12000, 20000, 24000, 72000),
    (2, "EPCM & Engineering Manager", "Juan Conde", 216000, 24000, 12000, 12000, 18000, 43200),
    (3, "Drilling & Well Services Manager", "Alexander Stulme", 180000, 24000, 12000, 12000, 18000, 36000),
    (4, "Operations & Maintenance Manager", "Félix Valderrama", 216000, 24000, 12000, 12000, 18000, 43200),
    (5, "Technical Manager (Geosciences)", "Alan McKeon", 180000, 24000, 12000, 12000, 18000, 36000),
    (6, "Reservoir Engineer", "JJI", 120000, 18000, 12000, 12000, 12000, 20000),
    (7, "Rig Company Man", "Marcelo Dantas", 180000, 24000, 12000, 12000, 18000, 36000),
    (8, "Planning & PMO", "Jose Miguel", 48000, 12000, 8000, 0, 7200, 12000),
    (9, "Operations Support", "Leticia Almeida", 48000, 12000, 8000, 12000, 7200, 8000),
]

SECONDEES = {"Alexander Stulme", "Alan McKeon"}

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
WHITE = "FFFFFF"
BLACK = "000000"
GRID = "B7B7B7"
CURRENCY = '$#,##0;[Red]($#,##0);-'
MONTHLY_CURRENCY = '$#,##0.00;[Red]($#,##0.00);-'


def build():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    wb = Workbook()
    ws = wb.active
    ws.title = "Annual Package"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "D2"

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(1, col, header)
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            left=Side(style="thin", color=BLACK),
            right=Side(style="thin", color=BLACK),
            top=Side(style="thin", color=BLACK),
            bottom=Side(style="medium", color=BLACK),
        )
    ws["A1"].comment = Comment(
        f"Transcribed from source image: {SOURCE.name}. Values verified in ordered overlapping crops.",
        "Manus",
    )
    ws.row_dimensions[1].height = 42

    for row_index, values in enumerate(ROWS, start=2):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row_index, col, value)
            cell.font = Font(name="Calibri", size=10, color=BLACK)
            cell.alignment = Alignment(
                horizontal="right" if col == 1 or col >= 4 else "left",
                vertical="center",
            )
            if col >= 4:
                cell.number_format = CURRENCY
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F5F9FC")
        total_cell = ws.cell(row_index, 10, f"=SUM(D{row_index}:I{row_index})")
        total_cell.number_format = CURRENCY
        total_cell.font = Font(name="Calibri", size=10, bold=True, color=BLACK)
        total_cell.alignment = Alignment(horizontal="right", vertical="center")
        if row_index % 2 == 0:
            total_cell.fill = PatternFill("solid", fgColor="F5F9FC")
        status_cell = ws.cell(
            row_index,
            11,
            "Secondee" if values[2] in SECONDEES else "Not Secondee",
        )
        status_cell.font = Font(name="Calibri", size=10, color=BLACK)
        status_cell.alignment = Alignment(horizontal="center", vertical="center")
        if row_index % 2 == 0:
            status_cell.fill = PatternFill("solid", fgColor="F5F9FC")
        for col in range(1, 12):
            ws.cell(row_index, col).border = Border(
                left=Side(style="thin", color=GRID),
                right=Side(style="thin", color=GRID),
                bottom=Side(style="thin", color=GRID),
            )
        ws.row_dimensions[row_index].height = 21

    total_row = 11
    ws.cell(total_row, 2, "TOTAL ANNUAL COST (9 positions)")
    ws.cell(total_row, 2).font = Font(name="Calibri", size=10, bold=True, color=BLACK)
    ws.cell(total_row, 2).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(total_row, 10, "=SUM(J2:J10)")
    ws.cell(total_row, 10).number_format = CURRENCY
    ws.cell(total_row, 10).font = Font(name="Calibri", size=10, bold=True, color=BLACK)
    ws.cell(total_row, 10).alignment = Alignment(horizontal="right", vertical="center")
    for col in range(1, 12):
        ws.cell(total_row, col).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(total_row, col).border = Border(
            top=Side(style="medium", color=BLACK),
            bottom=Side(style="double", color=BLACK),
        )
    ws.row_dimensions[total_row].height = 22

    table = Table(displayName="AnnualCompensationTable", ref="A1:K10")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    ws.add_table(table)

    widths = {
        "A": 5,
        "B": 36,
        "C": 24,
        "D": 20,
        "E": 18,
        "F": 18,
        "G": 23,
        "H": 18,
        "I": 16,
        "J": 22,
        "K": 18,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.auto_filter.ref = "A1:K10"
    ws.print_area = "A1:K11"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Annual Package | Page &P of &N"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = "666666"

    monthly = wb.create_sheet("Monthly Package")
    monthly.sheet_view.showGridLines = False
    monthly.freeze_panes = "D2"
    monthly_headers = [
        "#",
        "Position",
        "Person",
        "Base Salary (Monthly)",
        "Housing (Monthly)",
        "Medical Insurance (Monthly)",
        "Home Leave (Monthly)",
        "Company Vehicle (Monthly)",
        "Tax & Social (Monthly)",
        "TOTAL Monthly Package",
        "Secondment Status",
    ]
    for col, header in enumerate(monthly_headers, start=1):
        cell = monthly.cell(1, col, header)
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            left=Side(style="thin", color=BLACK),
            right=Side(style="thin", color=BLACK),
            top=Side(style="thin", color=BLACK),
            bottom=Side(style="medium", color=BLACK),
        )
    monthly["A1"].comment = Comment(
        "Monthly values are calculated as annual component divided by 12.",
        "Manus",
    )
    monthly.row_dimensions[1].height = 42

    for row_index in range(2, 11):
        for col in range(1, 4):
            annual_ref = ws.cell(row_index, col).coordinate
            cell = monthly.cell(row_index, col, f"='Annual Package'!{annual_ref}")
            cell.font = Font(name="Calibri", size=10, color=BLACK)
            cell.alignment = Alignment(
                horizontal="right" if col == 1 else "left",
                vertical="center",
            )
        for col in range(4, 10):
            annual_ref = ws.cell(row_index, col).coordinate
            cell = monthly.cell(row_index, col, f"='Annual Package'!{annual_ref}/12")
            cell.number_format = MONTHLY_CURRENCY
            cell.font = Font(name="Calibri", size=10, color=BLACK)
            cell.alignment = Alignment(horizontal="right", vertical="center")
        total_cell = monthly.cell(row_index, 10, f"=SUM(D{row_index}:I{row_index})")
        total_cell.number_format = MONTHLY_CURRENCY
        total_cell.font = Font(name="Calibri", size=10, bold=True, color=BLACK)
        total_cell.alignment = Alignment(horizontal="right", vertical="center")
        status_cell = monthly.cell(row_index, 11, f"='Annual Package'!K{row_index}")
        status_cell.font = Font(name="Calibri", size=10, color=BLACK)
        status_cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in range(1, 12):
            cell = monthly.cell(row_index, col)
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F5F9FC")
            cell.border = Border(
                left=Side(style="thin", color=GRID),
                right=Side(style="thin", color=GRID),
                bottom=Side(style="thin", color=GRID),
            )
        monthly.row_dimensions[row_index].height = 21

    monthly.cell(total_row, 2, "TOTAL MONTHLY COST (9 positions)")
    monthly.cell(total_row, 2).font = Font(name="Calibri", size=10, bold=True, color=BLACK)
    monthly.cell(total_row, 2).alignment = Alignment(horizontal="left", vertical="center")
    monthly.cell(total_row, 10, "=SUM(J2:J10)")
    monthly.cell(total_row, 10).number_format = MONTHLY_CURRENCY
    monthly.cell(total_row, 10).font = Font(name="Calibri", size=10, bold=True, color=BLACK)
    monthly.cell(total_row, 10).alignment = Alignment(horizontal="right", vertical="center")
    for col in range(1, 12):
        monthly.cell(total_row, col).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        monthly.cell(total_row, col).border = Border(
            top=Side(style="medium", color=BLACK),
            bottom=Side(style="double", color=BLACK),
        )
    monthly.row_dimensions[total_row].height = 22

    monthly_table = Table(displayName="MonthlyCompensationTable", ref="A1:K10")
    monthly_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    monthly.add_table(monthly_table)
    for column, width in widths.items():
        monthly.column_dimensions[column].width = width
    monthly.auto_filter.ref = "A1:K10"
    monthly.print_area = "A1:K11"
    monthly.print_title_rows = "1:1"
    monthly.page_setup.orientation = "landscape"
    monthly.page_setup.paperSize = monthly.PAPERSIZE_LETTER
    monthly.page_setup.fitToWidth = 1
    monthly.page_setup.fitToHeight = 1
    monthly.sheet_properties.pageSetUpPr.fitToPage = True
    monthly.oddFooter.center.text = "Monthly Package | Page &P of &N"
    monthly.oddFooter.center.size = 8
    monthly.oddFooter.center.color = "666666"

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUTPUT)
    print(OUTPUT)


def validate_values():
    wb = load_workbook(OUTPUT, data_only=False)
    ws = wb["Annual Package"]
    assert wb.sheetnames == ["Annual Package", "Monthly Package"]
    assert ws.max_row == 11 and ws.max_column == 11
    assert ws["B2"].value == "COO / PU General Manager"
    assert ws["C10"].value == "Leticia Almeida"
    assert ws["G9"].value == 0
    assert ws["G10"].value == 12000
    assert ws["J2"].value == "=SUM(D2:I2)"
    assert ws["J11"].value == "=SUM(J2:J10)"
    assert ws["K4"].value == "Secondee"
    assert ws["K6"].value == "Secondee"
    assert ws["K2"].value == "Not Secondee"
    monthly = wb["Monthly Package"]
    assert monthly["D2"].value == "='Annual Package'!D2/12"
    assert monthly["J2"].value == "=SUM(D2:I2)"
    assert monthly["J11"].value == "=SUM(J2:J10)"
    assert monthly["K4"].value == "='Annual Package'!K4"
    expected_totals = [524000, 325200, 282000, 325200, 282000, 194000, 282000, 87200, 95200]
    calculated = [sum(row[3:]) for row in ROWS]
    assert calculated == expected_totals
    assert sum(calculated) == 2396800
    print("SOURCE VALUES VALIDATED: $2,396,800")


if __name__ == "__main__":
    build()
    validate_values()
